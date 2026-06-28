#!/usr/bin/env python3
"""Generate the Final Presentation & Demo .pptx (Grilled Cheesin).

16:9, branded, 11 slides (title + F-1..F-10). Speaker notes carry the
'Say / defense' content. Import into the shared Google Slides deck via
File -> Import slides.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- palette ----------------------------------------------------------------
NAVY  = RGBColor(0x0B, 0x25, 0x45)
NAVY2 = RGBColor(0x13, 0x35, 0x5E)
TEAL  = RGBColor(0x2E, 0xC4, 0xB6)
AMBER = RGBColor(0xFF, 0x9F, 0x1C)
INK   = RGBColor(0x20, 0x2A, 0x37)
GRAY  = RGBColor(0x5A, 0x6B, 0x7B)
LIGHT = RGBColor(0xF2, 0xF5, 0xF9)
CARD  = RGBColor(0xE9, 0xEF, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _set_run(r, text, size, color, bold=False, italic=False, font=FONT):
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font


def add_textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    return tb, tf


def base_slide(title, tag):
    slide = prs.slides.add_slide(BLANK)
    # background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    _solid(bg, LIGHT)
    bg.shadow.inherit = False
    # title bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(1.02))
    _solid(bar, NAVY)
    bar.shadow.inherit = False
    # teal accent strip
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.02), SW, Inches(0.07))
    _solid(strip, TEAL)
    strip.shadow.inherit = False
    # title text
    tb, tf = add_textbox(slide, Inches(0.55), Inches(0.10), Inches(10.2), Inches(0.85),
                         anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    _set_run(p.add_run(), title, 26, WHITE, bold=True)
    # tag chip
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  SW - Inches(2.35), Inches(0.30), Inches(1.85), Inches(0.42))
    _solid(chip, TEAL)
    chip.shadow.inherit = False
    ctf = chip.text_frame
    ctf.word_wrap = True
    cp = ctf.paragraphs[0]
    cp.alignment = PP_ALIGN.CENTER
    _set_run(cp.add_run(), tag, 12, NAVY, bold=True)
    # footer
    ftb, ftf = add_textbox(slide, Inches(0.55), SH - Inches(0.42), Inches(12.2), Inches(0.32))
    fp = ftf.paragraphs[0]
    _set_run(fp.add_run(), "Grilled Cheesin  ·  Ocean Freight Forwarder Data Architecture  ·  MSDS 683 Final",
             10, GRAY)
    return slide


def add_bullets(slide, items, left, top, width, height, size=16, gap=6):
    """items: list of (text, level, kind) ; kind in {'bullet','head','note','code'}"""
    tb, tf = add_textbox(slide, left, top, width, height)
    first = True
    for it in items:
        text, level, kind = it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(gap)
        p.space_before = Pt(0)
        if kind == "head":
            _set_run(p.add_run(), text, size + 1, NAVY, bold=True)
        elif kind == "note":
            _set_run(p.add_run(), text, size - 2, GRAY, italic=True)
        elif kind == "code":
            _set_run(p.add_run(), text, size - 3, NAVY2, font=MONO)
        else:  # bullet
            bullet = "•  " if level == 0 else "–  "
            r = p.add_run()
            _set_run(r, bullet, size, TEAL, bold=True)
            # support simple **bold** segments
            for seg, bold in _split_bold(text):
                _set_run(p.add_run(), seg, size, INK, bold=bold)
    return tb


def _split_bold(text):
    out, buf, bold = [], "", False
    i = 0
    while i < len(text):
        if text[i:i+2] == "**":
            if buf:
                out.append((buf, bold))
                buf = ""
            bold = not bold
            i += 2
        else:
            buf += text[i]
            i += 1
    if buf:
        out.append((buf, bold))
    return out


def add_notes(slide, lines):
    notes = slide.notes_slide.notes_text_frame
    notes.text = lines[0]
    for ln in lines[1:]:
        p = notes.add_paragraph()
        p.text = ln


def add_chevrons(slide, labels, left, top, width, height, colors=None):
    n = len(labels)
    gap = Inches(0.06)
    cw = (width - gap * (n - 1)) / n
    colors = colors or [NAVY, NAVY2, TEAL, AMBER, NAVY][:n]
    for i, lab in enumerate(labels):
        x = left + i * (cw + gap)
        shp = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x, top, cw, height)
        _solid(shp, colors[i % len(colors)])
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        fontcol = NAVY if colors[i % len(colors)] in (TEAL, AMBER) else WHITE
        for j, part in enumerate(lab.split("\n")):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            _set_run(p.add_run(), part, 12 if j == 0 else 10,
                     fontcol, bold=(j == 0))


def add_table(slide, rows, left, top, width, col_widths=None, header=True,
              fontsize=12, row_h=0.34):
    nrows = len(rows)
    ncols = len(rows[0])
    height = Inches(row_h * nrows)
    gt = slide.shapes.add_table(nrows, ncols, left, top, width, height).table
    if col_widths:
        total = sum(col_widths)
        for c, w in enumerate(col_widths):
            gt.columns[c].width = Emu(int(width * w / total))
    # strip default banding style for cleaner look
    for r in range(nrows):
        gt.rows[r].height = Inches(row_h)
        for c in range(ncols):
            cell = gt.cell(r, c)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            is_head = header and r == 0
            cell.fill.solid()
            if is_head:
                cell.fill.fore_color.rgb = NAVY
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 else CARD
            para = cell.text_frame.paragraphs[0]
            cell.text_frame.word_wrap = True
            for seg, bold in _split_bold(str(rows[r][c])):
                _set_run(para.add_run(), seg, fontsize,
                         WHITE if is_head else INK, bold=bold or is_head)
    return gt


def add_card(slide, left, top, width, height, color=CARD):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _solid(card, color)
    card.shadow.inherit = False
    return card


# ============================================================ TITLE SLIDE ====
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
_solid(bg, NAVY); bg.shadow.inherit = False
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(4.55), SW, Inches(0.10))
_solid(band, TEAL); band.shadow.inherit = False
_, tf = add_textbox(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(2.2))
p = tf.paragraphs[0]
_set_run(p.add_run(), "Ocean Freight Forwarder", 44, WHITE, bold=True)
p2 = tf.add_paragraph()
_set_run(p2.add_run(), "Data Architecture — Final Presentation & Demo", 26, TEAL, bold=True)
_, tf2 = add_textbox(s, Inches(0.9), Inches(4.8), Inches(11.5), Inches(1.6))
p = tf2.paragraphs[0]
_set_run(p.add_run(), "Team Grilled Cheesin  —  P.J. Losiewicz · Borna Karimi · Alexander Mohun", 16, WHITE)
p = tf2.add_paragraph()
_set_run(p.add_run(), "A hybrid BigQuery (OLAP) + ArangoDB (network) architecture — the right store per workload", 14, RGBColor(0xC4, 0xD2, 0xE0), italic=True)
p = tf2.add_paragraph()
_set_run(p.add_run(), "MSDS 683 · Data Architecture", 13, GRAY)
add_notes(s, ["~10 min + Q&A. Spend the 5-minute core on the geo-fence event path (F-4/F-5) and edge weights → reroute (F-6), then land on the dashboard (F-7).",
              "Brief changed: Airflow is now BONUS; the demo must end in a consumable output (the web dashboard)."])

# ===================================================== F-1 RECAP =============
s = base_slide("Quick recap", "F-1 · ~30s")
add_bullets(s, [
    ("Team: Grilled Cheesin — P.J. Losiewicz · Borna Karimi · Alexander Mohun", 0, "bullet"),
    ("Domain: end-to-end data architecture for an **ocean freight forwarder / 3PL**", 0, "bullet"),
    ("Business question: which lanes, ports, and chokepoints put shipments at risk —", 0, "bullet"),
    ("and what does a disruption cost in transit time?", 1, "bullet"),
    ("Answer: a **hybrid** analytical layer — right store per workload:", 0, "bullet"),
    ("**BigQuery star warehouse** → OLAP questions (UC1 ETA reliability, UC2 congestion)", 1, "bullet"),
    ("**ArangoDB property graph** → network questions (UC3 chokepoints, UC4 rerouting)", 1, "bullet"),
    ("Shared business keys make the two stores one architecture, not two databases.", 0, "bullet"),
], Inches(0.6), Inches(1.45), Inches(12.1), Inches(5.2), size=17)
add_notes(s, ["The four questions don't all want the same database: roll-ups go to the columnar warehouse, reachability / shortest-path go to the graph.",
              "Audience already knows the domain — move fast."])

# ===================================================== F-2 SCHEMA ============
s = base_slide("Schema — what changed", "F-2 · ~30s")
add_bullets(s, [
    ("Core star + graph schema: **unchanged** since the midterm.", 0, "bullet"),
    ("Two date-partitioned, clustered facts (fact_voyage_leg, fact_port_call) + flat conformed dims; graph = 5 vertex + 4 edge collections.", 1, "bullet"),
    ("One refinement worth naming:", 0, "bullet"),
    ("added the **chokepoints vertex + transits_chokepoint edge** and a lane_key on every route edge,", 1, "bullet"),
    ("so a chokepoint closure resolves to the exact route edges that transit it.", 1, "bullet"),
    ("That turned UC3/UC4 from a diagram into a query that returns a number.", 1, "note"),
    ("Star over snowflake — still the defended call (columnar: storage cheap, joins are the cost).", 0, "bullet"),
], Inches(0.6), Inches(1.45), Inches(12.1), Inches(5.2), size=17)
add_notes(s, ["Schema is stable. The only meaningful change since the pitch is the chokepoint sub-graph, which made UC3/UC4 executable.",
              "Star-vs-snowflake rationale: m2-star-vs-snowflake.md."])

# ===================================================== F-3 PIPELINE ==========
s = base_slide("The pipeline, end to end", "F-3 · ~45s")
add_chevrons(s, ["RAW SOURCES\nAIS · ref · priors · synthetic",
                 "GCS BRONZE\nraw, immutable\nParquet/JSONL",
                 "GCS SILVER\nconform + DERIVE\ngeofence · weights",
                 "GOLD\nBigQuery star +\nArangoDB graph",
                 "PRODUCT\nNext.js dashboard"],
             Inches(0.5), Inches(1.4), Inches(12.3), Inches(1.15))
add_bullets(s, [
    ("Today we walk **one real path** end to end: raw AIS positions → geo-fenced port-calls & voyage-legs → both Gold stores → the dashboard.", 0, "bullet"),
    ("Representative subset, not full volume: Q1 2024, 4 US ports, ~1.88M AIS position rows.", 1, "note"),
    ("**One transform, two sinks:** both Gold stores read the *same* conformed Silver — never each other —", 0, "bullet"),
    ("so a shared-key metric reconciles between stores by construction.", 1, "bullet"),
    ("The hard parts, next: (1) turning lat/lon pings into **events**, (2) turning real trade indices into **edge weights**.", 0, "bullet"),
], Inches(0.6), Inches(2.95), Inches(12.1), Inches(3.6), size=16)
add_notes(s, ["I'll show the genuinely hard parts: geofencing pings into port-call events, and conditioning real trade indices into route edge weights. Then point the dashboard at the result."])

# ===================================================== F-4 GEOFENCE ==========
s = base_slide("Transform ① — geo-fencing raw AIS into events", "F-4 · ~1.5m · CORE")
add_bullets(s, [
    ("Input (Bronze): AIS fixes = (resolved IMO, WKB point, timestamp). We never trust the AIS free-text 'destination'.", 0, "bullet"),
    ("Geometry: a port fence is a **circle** — in-fence iff haversine_nm(fix, port) ≤ radius. (~10 lines of stdlib math; <0.5% error.)", 0, "bullet"),
    ("Event rule (per-vessel state machine, in time order): a port call = enter a fence + dwell ≥ a min threshold.", 0, "bullet"),
    ("**arrival** = first in-fence fix · **departure** = last in-fence fix before a *sustained* exit", 1, "bullet"),
    ("**debounce (2 fixes):** one stray out-ping between in-pings is still 'inside' → boundary jitter = one call", 1, "bullet"),
    ("**fence switch:** a fix in a *different* fence can't be jitter → close A, open B immediately", 1, "bullet"),
    ("**re-entry coalesce (≤12h):** drift out & back into the same fence reopens the same call", 1, "bullet"),
], Inches(0.6), Inches(1.4), Inches(12.1), Inches(3.7), size=15)
add_card(s, Inches(0.6), Inches(5.25), Inches(12.1), Inches(1.45))
add_bullets(s, [
    ("Tunable, documented defaults:  radius = 5 nm  ·  min-dwell = 1 h  ·  debounce = 2 fixes  ·  re-entry gap = 12 h", 0, "head"),
    ("Outside → enter fence (open) → dwell → out 2× (close) → emit if dwelled.   Source: silver/geofence.py + haversine.py", 0, "code"),
], Inches(0.78), Inches(5.38), Inches(11.7), Inches(1.2), size=14)
add_notes(s, ["Why a state machine, not GROUP BY port? An event is temporal — entry, continuous dwell, sustained exit. A naive group-by splits one messy call into five and invents zero-distance legs.",
              "Why circles not polygons? Centroid + radius = one cheap haversine per fix; harbour polygons need a geometry engine for <0.5% gain at this scale."])

# ===================================================== F-5 FACTS =============
s = base_slide("Transform ② — events → the two real facts", "F-5 · ~1m · CORE")
add_bullets(s, [
    ("silver/derive.py — pure, offline-testable, no I/O:", 0, "head"),
    ("derive_fact_port_calls → one fact_port_call per call; attaches conformed dim_port centroid (not AIS text), partitions by arrival date, provenance='real'.", 0, "bullet"),
    ("derive_voyage_legs → pairs each vessel's consecutive calls A→B into one fact_voyage_leg.", 0, "bullet"),
    ("Edge-case discipline = business rules: single-call vessel → 0 legs; same-port consecutive calls (origin==dest) → excluded (re-entry, not a transit).", 0, "bullet"),
], Inches(0.6), Inches(1.4), Inches(12.1), Inches(2.2), size=15)
add_table(s, [
    ["fact_voyage_leg (grain: one inter-port transit)", "source"],
    ["vessel_imo · origin_unlocode · dest_unlocode", "resolved keys"],
    ["transit_hours", "B.arrival − A.departure"],
    ["distance_nm", "haversine over port centroids (great-circle)"],
    ["schedule_delta", "actual − proforma  (NULL if no matching lane)"],
    ["dt (partition)", "origin-departure date"],
    ["provenance", "real"],
], Inches(0.6), Inches(3.6), Inches(12.1), col_widths=[7, 5], fontsize=13, row_h=0.40)
add_notes(s, ["This is where the 15 transformation points live: facts defined by business logic — what a port call is, what a voyage leg is, when a schedule delta is even meaningful — all in version-controlled, unit-tested pure functions (tests/test_geofence.py, test_derive.py).",
              "schedule_delta is NULL where no synthetic proforma lane matches — the honest real/synthetic seam; we don't fabricate a schedule."])

# ===================================================== F-6 EDGE WEIGHTS ======
s = base_slide("Transform ③ — the edge-weight logic", "F-6 · ~1.5m · CORE")
add_bullets(s, [
    ("Route edge weights are **derived from real trade indices** (no free bilateral port-pair feed exists — we condition, never invent):", 0, "head"),
    ("Lane plausibility:  lane_weight(A,B) = norm(LSCI[A]) × norm(LSCI[B]) × norm(Comtrade[A→B])", 0, "code"),
    ("Reliability/delay:  expected_delay(c) = 72h × (max_LPI − LPI[c]) / (max_LPI − min_LPI)  → seeded lognormal", 0, "code"),
    ("Four route weights: distance_nm (haversine) · transit_time_hours = distance_nm / 18 kn · service_frequency (LSCI) · reliability/expected_delay (LPI)", 0, "bullet"),
    ("Provenance on weights too: 'real' only when overlaid from observed fact_voyage_leg (US→US); else 'synthetic'.", 0, "bullet"),
], Inches(0.6), Inches(1.4), Inches(12.1), Inches(2.9), size=14)
add_bullets(s, [("Then the weight does work — UC4 reroute (K_SHORTEST_PATHS, weighted by transit_time_hours):", 0, "head")],
            Inches(0.6), Inches(4.35), Inches(12.1), Inches(0.4), size=15)
add_table(s, [
    ["Scenario", "Path", "Hours"],
    ["Baseline  USNYC → CNSHA", "direct", "**355.97 h**"],
    ["SUEZ closed (disable transiting lanes, re-run)", "USNYC → **USLAX** → CNSHA", "**432.19 h**"],
    ["Reroute cost", "detours via USLAX", "**+76.22 h**"],
], Inches(0.6), Inches(4.85), Inches(12.1), col_widths=[6, 5, 2.2], fontsize=13, row_h=0.42)
add_notes(s, ["'Where's your route table?' — here. No magic bilateral lane feed exists, so we condition synthetic lanes on three real indices (LSCI, Comtrade, LPI) and document them as priors, never promoted to facts.",
              "Why K_SHORTEST_PATHS not SHORTEST_PATH? A bare shortest-path filters disabled lanes AFTER finding the optimum, so closing the best lane returns empty. K-shortest enumerates by weight; we keep the cheapest path avoiding the closed chokepoint → a genuine +delta detour."])

# ===================================================== F-7 END PRODUCT =======
s = base_slide("End product — the dashboard", "F-7 · ~1m · OUTPUT")
add_bullets(s, [
    ("web/ — Next.js 16, deployed on Vercel. The brief's required 'concrete, consumable output'.", 0, "head"),
    ("Four pages /uc1–/uc4: UC1/UC2 = KPI tables + trend charts (recharts) off the BigQuery star;", 0, "bullet"),
    ("UC3/UC4 = interactive map (deck.gl + MapLibre) — chokepoints, transit share, live reroute path on closure.", 0, "bullet"),
    ("Pointed straight at Gold: server components read live BigQuery + ArangoDB; with no creds they fall back to frozen goldens —", 0, "bullet"),
    ("so the demo renders from a clean clone and **cannot fail live**. A recorded backup is the belt-and-suspenders.", 1, "bullet"),
], Inches(0.6), Inches(1.45), Inches(12.1), Inches(3.1), size=16)
add_card(s, Inches(0.6), Inches(4.8), Inches(12.1), Inches(1.25), color=NAVY)
_, tf = add_textbox(s, Inches(0.85), Inches(4.95), Inches(11.6), Inches(1.0), anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]
_set_run(p.add_run(), "The headline a planner reads off the screen:", 14, TEAL, bold=True)
p = tf.add_paragraph()
_set_run(p.add_run(), "“Closing Suez adds +76 hours to New York → Shanghai, rerouting via Los Angeles.”", 18, WHITE, bold=True)
add_notes(s, ["This is the consumable output — the gold layer rendered as a map, not rows in a table. Show the UC3/UC4 page; toggle a closure live.",
              "Live toggle is a 'look, it's real' aside; never required to render."])

# ===================================================== F-8 TECH STACK ========
s = base_slide("Tech stack", "F-8 · 1–2m")
add_table(s, [
    ["Layer", "Tool", "Role"],
    ["Raw + staging", "**GCS**", "immutable Bronze + conformed Silver (Parquet/JSONL, dt= partitioned)"],
    ["OLAP store", "**BigQuery**", "star warehouse — UC1/UC2 roll-ups (partition + cluster)"],
    ["Graph store", "**ArangoDB** (managed)", "ocean_network graph — UC3/UC4 traversal + K_SHORTEST_PATHS"],
    ["Orchestration (bonus)", "**Airflow 3.0**", "ofa_warehouse DAG: conform → load_bq ∥ load_arango → verify"],
    ["Transforms / glue", "**Python 3.11**", "conform / derive / geofence / conditioning"],
    ["Tabular + format", "**pandas + pyarrow**", "shape dims/facts; columnar Parquet, JSONL for events"],
    ["Synthetic data", "**NumPy + Faker**", "seeded, reproducible delays/volumes/identifiers"],
    ["Graph driver", "**python-arango**", "idempotent UPSERT loader from the same Silver"],
    ["End product", "**Next.js + deck.gl + Vercel**", "the consumable dashboard over live BQ/Arango"],
], Inches(0.5), Inches(1.35), Inches(12.35), col_widths=[3.0, 3.0, 7.2], fontsize=12, row_h=0.42)
add_bullets(s, [("Per-tool justification + alternatives (for the 'why this / what alternatives' follow-up): docs/deck/tech-stack/", 0, "note")],
            Inches(0.6), Inches(6.75), Inches(12.1), Inches(0.4), size=12)
add_notes(s, ["The brief warns: I'll ask about one tool, why you picked it, what alternatives. Each tool has a file with a ready answer.",
              "Headline defenses: BigQuery (serverless columnar — star pays off, no cluster), ArangoDB (graph + team expertise/managed cluster), GCS (native landing both BQ & Composer read), star>snowflake (joins cost more than storage on columnar)."])

# ===================================================== F-9 BUDGET ============
s = base_slide("Budget & cost model", "F-9 · 1–2m")
add_bullets(s, [("Course-scale today: Q1 2024, 4 ports, ~1.88M AIS rows; graph ≤ ~100K edges. Live guard: **$50/mo GCP budget** w/ 50/90/100% alerts.", 0, "head")],
            Inches(0.6), Inches(1.35), Inches(12.1), Inches(0.6), size=14)
add_table(s, [
    ["Component", "Driver", "Course-scale", "Production (assumptions below)"],
    ["GCS storage", "bytes landed", "<1 GB → ~$0.02/mo", "2 TB → ~$40/mo"],
    ["BigQuery storage", "table bytes", "<2 GB → ~$0.04/mo", "1 TB native → ~$20/mo"],
    ["BigQuery query", "bytes **scanned**", "~$0 (free tier)", "50 TB/mo @ $6.25/TB → ~$310/mo"],
    ["ArangoDB cluster", "instance hrs", "team cluster → ~$0", "small prod cluster → ~$200–400/mo"],
    ["Airflow (Composer 3)", "env hrs", "not always-on → ~$0", "always-on small env → ~$300–450/mo"],
    ["Web (Vercel)", "requests", "hobby → ~$0", "Pro → ~$20/mo"],
    ["**Total**", "", "**≈ $0–5/mo (under $50 cap)**", "**≈ $900–1,250/mo (~$11–15K/yr)**"],
], Inches(0.5), Inches(2.0), Inches(12.35), col_widths=[3.0, 2.4, 3.2, 4.6], fontsize=12, row_h=0.40)
add_bullets(s, [
    ("Production assumptions: full global AIS, daily batch loads, ~50 TB/mo scanned, always-on Composer + small Arango cluster.", 0, "note"),
    ("Biggest lever = **bytes scanned** → which is *why* we partition on dt and cluster on FKs. The physical design IS the cost model.", 0, "bullet"),
], Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.1), size=13)
add_notes(s, ["Today we're under $5/mo against a $50 cap. The production number is dominated by always-on compute (Composer + Arango), not storage. The graph stays cheap because we bound it to a defensible slice."])

# ===================================================== F-10 BONUS ============
s = base_slide("Bonus — orchestration with Airflow", "F-10 · +bonus")
add_chevrons(s, ["stage_conform\nbuild Silver",
                 "load_bigquery\nGCS → star",
                 "load_arango\nUPSERT graph",
                 "verify\nfan-in: counts +\ncross-store + UC"],
             Inches(0.6), Inches(1.45), Inches(12.1), Inches(1.15),
             colors=[NAVY, TEAL, TEAL, AMBER])
add_bullets(s, [
    ("Airflow is now bonus — and **we built it.** The ofa_warehouse **Airflow 3.0** DAG is the implemented cloud-ETL.", 0, "bullet"),
    ("One conform → **two parallel load legs** (load_bigquery ∥ load_arango) → **fan-in verify**", 0, "bullet"),
    ("(row counts + cross-store reconcile + live UC3/UC4 anti-degeneracy).", 1, "bullet"),
    ("Plain Apache Airflow, **Composer-portable**: standard operators only; runs locally via airflow dags test; lifts to Cloud Composer 3 unchanged.", 0, "bullet"),
    ("verify fans in on both legs → a half-loaded store can't pass the gate by accident.", 0, "bullet"),
], Inches(0.6), Inches(3.0), Inches(12.1), Inches(3.4), size=16)
add_notes(s, ["Per the brief, orchestration is bonus now — but it's done. One conform feeding two parallel loads is the 'one transform, two sinks' contract: that's what guarantees the warehouse and the graph are the same architecture."])

out = "docs/deck/final-presentation.pptx"
prs.save(out)
print("WROTE", out, "with", len(prs.slides._sldIdLst), "slides")
